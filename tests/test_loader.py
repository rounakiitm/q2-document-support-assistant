"""Loader/splitter tests: section paths must be citation-grade metadata."""

from pathlib import Path

from backend.ingest.loader import fmt_for, load_any

MD_DOC = """# Top Title

## Free Look
You may return the policy within fifteen days. Refunds apply.

### Refund Calculation
Refund equals premium less risk premium. Stamp duty is retained.

## Grace Period
Thirty days for annual mode. Fifteen for monthly.
"""


def test_fmt_dispatch():
    assert fmt_for("a.pdf") == "pdf"
    assert fmt_for("b.DOCX") == "docx"
    assert fmt_for("c.md") == "md"
    assert fmt_for("d.txt") == ""


def test_md_section_paths(tmp_path: Path):
    p = tmp_path / "doc.md"
    p.write_text(MD_DOC, encoding="utf-8")
    chunks = load_any(p, "md")
    assert chunks, "no chunks produced"
    paths = [c.metadata["section_path"] for c in chunks]
    assert any("Free Look" in x for x in paths)
    assert any("Refund Calculation" in x for x in paths)
    # every chunk belongs to exactly one section path
    for c in chunks:
        assert c.metadata["section_path"] in {"Top Title › Free Look › Refund Calculation", "Top Title › Free Look", "Top Title › Grace Period"} or "Free Look" in c.metadata["section_path"] or "Grace" in c.metadata["section_path"]


def test_docx_heading_paths(tmp_path: Path):
    from docx import Document

    p = tmp_path / "doc.docx"
    d = Document()
    d.add_paragraph("Escalation", style="Heading 1")
    d.add_paragraph("Level 1: branch contact. Acknowledged in 3 days.", style="Normal")
    d.add_paragraph("Ombudsman", style="Heading 2")
    d.add_paragraph("Award binds the company.", style="Normal")
    d.save(p)
    chunks = load_any(p, "docx")
    texts = [c.page_content for c in chunks]
    paths = [c.metadata["section_path"] for c in chunks]
    assert any("Award binds" in t for t in texts)
    assert any("Escalation" in x and "Ombudsman" in x for x in paths)


def test_pdf_pages(tmp_path: Path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 14)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 8, "Purpose")
    pdf.ln(1)
    pdf.set_font("helvetica", "", 10)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 5, "The profiler records suitability before the proposal is signed.")
    p = tmp_path / "doc.pdf"
    pdf.output(p)
    chunks = load_any(p, "pdf")
    assert chunks
    assert all(c.metadata["page"] == 1 for c in chunks)
    assert any("profiler records" in c.page_content.lower() for c in chunks)